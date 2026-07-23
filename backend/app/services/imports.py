import io

from openpyxl import load_workbook

QUESTION_HEADER_TEXT = "savol"
QUESTION_HEADER_CORRECT = "togri_javob"
QUESTION_HEADER_TOPIC = "mavzu"
QUESTION_HEADER_VARIANT_PREFIX = "variant_"


def _parse_correct_index(raw, option_count: int) -> int | None:
    if raw is None:
        return None
    s = str(raw).strip().upper()
    if len(s) == 1 and "A" <= s <= "F":
        idx = ord(s) - ord("A")
    elif s.isdigit():
        idx = int(s) - 1
    else:
        return None
    return idx if 0 <= idx < option_count else None


def parse_questions_xlsx(file_bytes: bytes) -> dict:
    """Parse an xlsx question bank per spec §7.1.

    Returns {"questions": [{"text", "options", "correct_index", "topic"}], "errors": [str, ...]}.
    """
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    if not rows:
        return {"questions": [], "errors": ["Fayl bo'sh"]}

    header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]

    def col_index(name: str) -> int | None:
        return header.index(name) if name in header else None

    idx_text = col_index(QUESTION_HEADER_TEXT)
    idx_correct = col_index(QUESTION_HEADER_CORRECT)
    idx_topic = col_index(QUESTION_HEADER_TOPIC)
    variant_indices = sorted(i for i, h in enumerate(header) if h.startswith(QUESTION_HEADER_VARIANT_PREFIX))

    if idx_text is None or idx_correct is None or not variant_indices:
        return {
            "questions": [],
            "errors": [
                "Ustunlar topilmadi: 'savol', kamida bitta 'variant_*' va 'togri_javob' bo'lishi shart"
            ],
        }

    questions = []
    errors = []

    for row_num, row in enumerate(rows[1:], start=2):
        text_raw = row[idx_text] if idx_text < len(row) else None
        text = str(text_raw).strip() if text_raw is not None else ""
        if not text:
            errors.append(f"{row_num}-qator: savol matni bo'sh")
            continue

        options = []
        for vi in variant_indices:
            val = row[vi] if vi < len(row) else None
            if val is not None and str(val).strip():
                options.append(str(val).strip())

        if len(options) < 2:
            errors.append(f"{row_num}-qator: kamida 2 ta variant kerak")
            continue

        raw_correct = row[idx_correct] if idx_correct < len(row) else None
        correct_index = _parse_correct_index(raw_correct, len(options))
        if correct_index is None:
            errors.append(f"{row_num}-qator: togri_javob noto'g'ri qiymat ('{raw_correct}')")
            continue

        topic = None
        if idx_topic is not None and idx_topic < len(row) and row[idx_topic] is not None:
            topic_val = str(row[idx_topic]).strip()
            topic = topic_val or None

        questions.append(
            {
                "text": text,
                "options": options,
                "correct_index": correct_index,
                "topic": topic,
            }
        )

    return {"questions": questions, "errors": errors}
